from pathlib import Path
from math import atan2, cos, sin, pi

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "tmp" / "er_document"
BUILD.mkdir(parents=True, exist_ok=True)
DIAGRAM_PATH = BUILD / "easyattend_er_diagram.png"
DOCX_PATH = BUILD / "EasyAttend_Database_ER_Diagram_A4.docx"


def font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


W, H = 5200, 3150
BG = "#F8FAFC"
INK = "#162234"
MUTED = "#5F6B7A"
LINE = "#7C8DA5"

entities = {
    "USERS": ["PK  id", "UK  username", "password_hash", "full_name", "role", "status", "created_at"],
    "ACADEMIC_YEARS": ["PK  id", "UK  year_level", "UK  name"],
    "SEMESTERS": ["PK  id", "FK  academic_year_id", "semester_number", "name"],
    "CLASSES": ["PK  id", "FK  academic_year_id", "UK  name"],
    "STUDENTS": ["PK  id", "FK  user_id", "UK  student_no", "FK  class_id", "FK  semester_id", "year_level", "device_uuid"],
    "TEACHERS": ["PK  id", "FK  user_id", "UK  staff_no", "FK  subject_id", "year_level"],
    "SUBJECTS": ["PK  id", "code", "name", "year_level", "FK  semester_id", "registration_enabled"],
    "TEACHER_TERMS": ["PK  id", "FK  teacher_id", "FK  semester_id", "FK  class_id", "created_at"],
    "TEACHER_SUBJECTS": ["PK  id", "FK  teacher_id", "FK  teacher_term_id", "FK  subject_id", "created_at"],
    "CLASSROOMS": ["PK  id", "UK  name"],
    "SCHEDULES": ["PK  id", "FK  subject_id", "FK  teacher_id", "FK  classroom_id", "day_of_week", "start_time / end_time"],
    "ATTENDANCE_SESSIONS": ["PK  id", "FK  teacher_id", "FK  subject_id", "FK  teacher_subject_id", "FK  schedule_id", "title", "starts_at / ended_at", "active"],
    "API_TOKENS": ["PK  id", "FK  user_id", "UK  token", "expires_at", "created_at"],
    "NOTIFICATIONS": ["PK  id", "FK  user_id", "message", "read_at", "created_at"],
    "QR_CODES": ["PK  id", "FK  session_id", "UK  token_hash", "display_token", "expires_at", "active"],
    "ATTENDANCE": ["PK  id", "FK  session_id", "FK  student_id", "status", "latitude / longitude", "accuracy", "distance_from_classroom", "recorded_at"],
}

positions = {
    "USERS": (100, 90), "ACADEMIC_YEARS": (1375, 90), "SEMESTERS": (2650, 90), "CLASSES": (3925, 90),
    "STUDENTS": (100, 850), "TEACHERS": (1375, 850), "SUBJECTS": (2650, 850), "TEACHER_TERMS": (3925, 850),
    "TEACHER_SUBJECTS": (100, 1610), "CLASSROOMS": (1375, 1610), "SCHEDULES": (2650, 1610), "ATTENDANCE_SESSIONS": (3925, 1610),
    "API_TOKENS": (100, 2370), "NOTIFICATIONS": (1375, 2370), "QR_CODES": (2650, 2370), "ATTENDANCE": (3925, 2370),
}

relationships = [
    ("STUDENTS", "USERS"), ("TEACHERS", "USERS"), ("API_TOKENS", "USERS"), ("NOTIFICATIONS", "USERS"),
    ("SEMESTERS", "ACADEMIC_YEARS"), ("CLASSES", "ACADEMIC_YEARS"),
    ("SUBJECTS", "SEMESTERS"), ("STUDENTS", "SEMESTERS"), ("STUDENTS", "CLASSES"),
    ("TEACHERS", "SUBJECTS"), ("TEACHER_TERMS", "TEACHERS"), ("TEACHER_TERMS", "SEMESTERS"), ("TEACHER_TERMS", "CLASSES"),
    ("TEACHER_SUBJECTS", "TEACHERS"), ("TEACHER_SUBJECTS", "TEACHER_TERMS"), ("TEACHER_SUBJECTS", "SUBJECTS"),
    ("SCHEDULES", "SUBJECTS"), ("SCHEDULES", "TEACHERS"), ("SCHEDULES", "CLASSROOMS"),
    ("ATTENDANCE_SESSIONS", "TEACHERS"), ("ATTENDANCE_SESSIONS", "SUBJECTS"),
    ("ATTENDANCE_SESSIONS", "TEACHER_SUBJECTS"), ("ATTENDANCE_SESSIONS", "SCHEDULES"),
    ("QR_CODES", "ATTENDANCE_SESSIONS"), ("ATTENDANCE", "ATTENDANCE_SESSIONS"), ("ATTENDANCE", "STUDENTS"),
]

box_w, box_h = 1120, 650


def box_center(name):
    x, y = positions[name]
    return x + box_w / 2, y + box_h / 2


def boundary_point(name, toward):
    cx, cy = box_center(name)
    tx, ty = toward
    dx, dy = tx - cx, ty - cy
    if not dx and not dy:
        return cx, cy
    scale = min((box_w / 2) / abs(dx) if dx else 1e9, (box_h / 2) / abs(dy) if dy else 1e9)
    return cx + dx * scale, cy + dy * scale


def arrow(draw, child, parent):
    pc = box_center(parent)
    cc = box_center(child)
    start = boundary_point(child, pc)
    end = boundary_point(parent, cc)
    draw.line([start, end], fill=LINE, width=8)
    angle = atan2(end[1] - start[1], end[0] - start[0])
    length = 30
    spread = pi / 7
    a = (end[0] - length * cos(angle - spread), end[1] - length * sin(angle - spread))
    b = (end[0] - length * cos(angle + spread), end[1] - length * sin(angle + spread))
    draw.polygon([end, a, b], fill=LINE)


def draw_diagram():
    image = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(image)
    for child, parent in relationships:
        arrow(draw, child, parent)

    header_font = font(52, True)
    body_font = font(43)
    key_font = font(43, True)
    palette = {
        "identity": ("#DDEEFF", "#2F6FAE"),
        "academic": ("#E6F4E8", "#3D7A4B"),
        "assignment": ("#FFF2CC", "#A87900"),
        "attendance": ("#EADFF4", "#744A91"),
        "support": ("#ECEFF3", "#5E6B78"),
    }
    category = {
        "USERS": "identity", "STUDENTS": "identity", "TEACHERS": "identity",
        "ACADEMIC_YEARS": "academic", "SEMESTERS": "academic", "CLASSES": "academic", "SUBJECTS": "academic",
        "TEACHER_TERMS": "assignment", "TEACHER_SUBJECTS": "assignment", "CLASSROOMS": "assignment", "SCHEDULES": "assignment",
        "ATTENDANCE_SESSIONS": "attendance", "QR_CODES": "attendance", "ATTENDANCE": "attendance",
        "API_TOKENS": "support", "NOTIFICATIONS": "support",
    }

    for name, fields in entities.items():
        x, y = positions[name]
        fill, border = palette[category[name]]
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=24, fill="white", outline=border, width=8)
        draw.rounded_rectangle((x, y, x + box_w, y + 105), radius=24, fill=fill, outline=border, width=8)
        draw.rectangle((x + 4, y + 80, x + box_w - 4, y + 110), fill=fill)
        bbox = draw.textbbox((0, 0), name, font=header_font)
        draw.text((x + (box_w - (bbox[2] - bbox[0])) / 2, y + 24), name, font=header_font, fill=INK)
        line_y = y + 126
        line_gap = min(66, 500 // max(1, len(fields)))
        for field in fields:
            is_key = field.startswith(("PK", "FK", "UK"))
            draw.text((x + 42, line_y), field, font=key_font if is_key else body_font, fill=INK if is_key else MUTED)
            line_y += line_gap

    legend_font = font(32)
    draw.text((110, 3078), "Arrow direction: foreign key table  →  referenced table     PK primary key     FK foreign key     UK unique key", font=legend_font, fill=MUTED)
    image.save(DIAGRAM_PATH, quality=95, dpi=(420, 420))


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def remove_paragraph_borders(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(7)
    section.left_margin = Mm(10)
    section.right_margin = Mm(10)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(31, 41, 55)

    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(20)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    remove_paragraph_borders(title_style)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing = 1.15
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title.add_run("EasyAttend Database Entity Relationship Diagram")
    remove_paragraph_borders(title)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_after = Pt(5)
    intro.add_run(
        "This diagram presents the principal entities and foreign key relationships used for user management, academic assignments, QR sessions, GPS verified attendance, and supporting records."
    )

    picture_p = doc.add_paragraph()
    picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_p.paragraph_format.space_after = Pt(2)
    picture_p.add_run().add_picture(str(DIAGRAM_PATH), width=Inches(10.85))

    props = doc.core_properties
    props.title = "EasyAttend Database Entity Relationship Diagram"
    props.subject = "A4 database ER diagram"
    props.author = "EasyAttend Project Team"
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    draw_diagram()
    build_docx()
    print(DOCX_PATH)
