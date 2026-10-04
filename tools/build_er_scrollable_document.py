from pathlib import Path
from math import atan2, cos, sin, pi

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "tmp" / "er_scrollable_document"
BUILD.mkdir(parents=True, exist_ok=True)
DOCX_PATH = BUILD / "EasyAttend_Database_ER_Diagrams_Clear_Relationships_A4.docx"

W, H = 4800, 2700
BOX_W, BOX_H = 1080, 650
BG = "#F7F9FC"
INK = "#152235"
MUTED = "#5E6A78"
RELATION = "#496A92"


def get_font(size, bold=False):
    names = ["arialbd.ttf", "calibrib.ttf"] if bold else ["arial.ttf", "calibri.ttf"]
    for name in names:
        path = Path("C:/Windows/Fonts") / name
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


HEADER_FONT = get_font(54, True)
BODY_FONT = get_font(48)
KEY_FONT = get_font(48, True)
LABEL_FONT = get_font(38, True)


PALETTE = {
    "identity": ("#DDEEFF", "#2F6FAE"),
    "academic": ("#E6F4E8", "#3D7A4B"),
    "assignment": ("#FFF2CC", "#A87900"),
    "attendance": ("#EADFF4", "#744A91"),
    "support": ("#ECEFF3", "#5E6B78"),
}


ENTITY_STYLE = {
    "USERS": "identity", "STUDENTS": "identity", "TEACHERS": "identity",
    "ACADEMIC_YEARS": "academic", "SEMESTERS": "academic", "CLASSES": "academic", "SUBJECTS": "academic",
    "TEACHER_TERMS": "assignment", "TEACHER_SUBJECTS": "assignment", "CLASSROOMS": "assignment", "SCHEDULES": "assignment",
    "ATTENDANCE_SESSIONS": "attendance", "QR_CODES": "attendance", "ATTENDANCE": "attendance",
    "API_TOKENS": "support", "NOTIFICATIONS": "support",
}


FIELDS = {
    "USERS": ["PK  id", "UK  username", "password_hash", "full_name", "role", "status", "created_at"],
    "STUDENTS": ["PK  id", "FK  user_id", "UK  student_no", "FK  class_id", "FK  semester_id", "year_level", "device_uuid"],
    "TEACHERS": ["PK  id", "FK  user_id", "UK  staff_no", "FK  subject_id", "year_level"],
    "ACADEMIC_YEARS": ["PK  id", "UK  year_level", "UK  name"],
    "SEMESTERS": ["PK  id", "FK  academic_year_id", "semester_number", "name"],
    "CLASSES": ["PK  id", "FK  academic_year_id", "UK  name"],
    "SUBJECTS": ["PK  id", "code", "name", "year_level", "FK  semester_id", "registration_enabled"],
    "TEACHER_TERMS": ["PK  id", "FK  teacher_id", "FK  semester_id", "FK  class_id", "created_at"],
    "TEACHER_SUBJECTS": ["PK  id", "FK  teacher_id", "FK  teacher_term_id", "FK  subject_id", "created_at"],
    "CLASSROOMS": ["PK  id", "UK  name"],
    "SCHEDULES": ["PK  id", "FK  subject_id", "FK  teacher_id", "FK  classroom_id", "day_of_week", "start_time", "end_time"],
    "ATTENDANCE_SESSIONS": ["PK  id", "FK  teacher_id", "FK  subject_id", "FK  teacher_subject_id", "FK  schedule_id", "title", "starts_at / ended_at", "active"],
    "QR_CODES": ["PK  id", "FK  session_id", "UK  token_hash", "display_token", "expires_at", "active"],
    "ATTENDANCE": ["PK  id", "FK  session_id", "FK  student_id", "status", "latitude / longitude", "accuracy", "distance_from_classroom", "recorded_at"],
    "NOTIFICATIONS": ["PK  id", "FK  user_id", "message", "read_at", "created_at"],
    "API_TOKENS": ["PK  id", "FK  user_id", "UK  token", "expires_at", "created_at"],
}


RELATIONSHIPS = [
    ("USERS", "1", "STUDENTS", "0..1", "students.user_id", "A user may have one student profile"),
    ("USERS", "1", "TEACHERS", "0..1", "teachers.user_id", "A user may have one teacher profile"),
    ("USERS", "1", "NOTIFICATIONS", "0..*", "notifications.user_id", "A user receives notifications"),
    ("USERS", "1", "API_TOKENS", "0..*", "api_tokens.user_id", "A user owns API login tokens"),
    ("ACADEMIC_YEARS", "1", "SEMESTERS", "0..*", "semesters.academic_year_id", "An academic year contains semesters"),
    ("ACADEMIC_YEARS", "1", "CLASSES", "0..*", "classes.academic_year_id", "An academic year contains classes"),
    ("SEMESTERS", "0..1", "SUBJECTS", "0..*", "subjects.semester_id", "A semester contains subjects"),
    ("SEMESTERS", "0..1", "STUDENTS", "0..*", "students.semester_id", "Students are assigned to a semester"),
    ("CLASSES", "0..1", "STUDENTS", "0..*", "students.class_id", "A class contains students"),
    ("SUBJECTS", "0..1", "TEACHERS", "0..*", "teachers.subject_id", "A teacher may retain a primary subject"),
    ("TEACHERS", "1", "TEACHER_TERMS", "0..*", "teacher_terms.teacher_id", "A teacher receives term assignments"),
    ("SEMESTERS", "1", "TEACHER_TERMS", "0..*", "teacher_terms.semester_id", "A term assignment belongs to a semester"),
    ("CLASSES", "1", "TEACHER_TERMS", "0..*", "teacher_terms.class_id", "A term assignment belongs to a class"),
    ("TEACHERS", "1", "TEACHER_SUBJECTS", "0..*", "teacher_subjects.teacher_id", "A teacher teaches assigned subjects"),
    ("TEACHER_TERMS", "0..1", "TEACHER_SUBJECTS", "0..*", "teacher_subjects.teacher_term_id", "A term assignment contains subject assignments"),
    ("SUBJECTS", "1", "TEACHER_SUBJECTS", "0..*", "teacher_subjects.subject_id", "A subject appears in teaching assignments"),
    ("SUBJECTS", "1", "SCHEDULES", "0..*", "schedules.subject_id", "A subject has scheduled periods"),
    ("TEACHERS", "1", "SCHEDULES", "0..*", "schedules.teacher_id", "A teacher follows scheduled periods"),
    ("CLASSROOMS", "0..1", "SCHEDULES", "0..*", "schedules.classroom_id", "A classroom hosts scheduled periods"),
    ("TEACHERS", "1", "ATTENDANCE_SESSIONS", "0..*", "attendance_sessions.teacher_id", "A teacher creates attendance sessions"),
    ("SUBJECTS", "1", "ATTENDANCE_SESSIONS", "0..*", "attendance_sessions.subject_id", "A session records attendance for a subject"),
    ("TEACHER_SUBJECTS", "0..1", "ATTENDANCE_SESSIONS", "0..*", "attendance_sessions.teacher_subject_id", "An assignment authorizes a session"),
    ("SCHEDULES", "0..1", "ATTENDANCE_SESSIONS", "0..*", "attendance_sessions.schedule_id", "A schedule may start attendance sessions"),
    ("ATTENDANCE_SESSIONS", "1", "QR_CODES", "0..*", "qr_codes.session_id", "A session generates QR codes"),
    ("ATTENDANCE_SESSIONS", "1", "ATTENDANCE", "0..*", "attendance.session_id", "A session contains attendance records"),
    ("STUDENTS", "1", "ATTENDANCE", "0..*", "attendance.student_id", "A student submits attendance records"),
]


def center(pos):
    return pos[0] + BOX_W / 2, pos[1] + BOX_H / 2


def boundary(pos, toward):
    cx, cy = center(pos)
    dx, dy = toward[0] - cx, toward[1] - cy
    scale = min((BOX_W / 2) / abs(dx) if dx else 1e9, (BOX_H / 2) / abs(dy) if dy else 1e9)
    return cx + dx * scale, cy + dy * scale


def draw_arrow_head(draw, previous, end):
    angle = atan2(end[1] - previous[1], end[0] - previous[0])
    length = 34
    spread = pi / 7
    a = (end[0] - length * cos(angle - spread), end[1] - length * sin(angle - spread))
    b = (end[0] - length * cos(angle + spread), end[1] - length * sin(angle + spread))
    draw.polygon([end, a, b], fill=RELATION)


def draw_relationship(draw, positions, parent, child, label, route=None):
    parent_center = center(positions[parent])
    child_center = center(positions[child])
    first_target = route[0] if route else child_center
    last_source = route[-1] if route else parent_center
    points = [boundary(positions[parent], first_target)] + (route or []) + [boundary(positions[child], last_source)]
    draw.line(points, fill=RELATION, width=9, joint="curve")
    draw_arrow_head(draw, points[-2], points[-1])
    segment_lengths = []
    total = 0
    for a, b in zip(points, points[1:]):
        length = ((b[0] - a[0]) ** 2 + (b[1] - a[1]) ** 2) ** 0.5
        segment_lengths.append(length)
        total += length
    target = total / 2
    walked = 0
    lx, ly = points[0]
    for (a, b), length in zip(zip(points, points[1:]), segment_lengths):
        if walked + length >= target:
            ratio = (target - walked) / max(1, length)
            lx = a[0] + (b[0] - a[0]) * ratio
            ly = a[1] + (b[1] - a[1]) * ratio
            break
        walked += length
    bbox = draw.textbbox((0, 0), label, font=LABEL_FONT)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    pad_x, pad_y = 18, 10
    draw.rounded_rectangle((lx - tw / 2 - pad_x, ly - th / 2 - pad_y, lx + tw / 2 + pad_x, ly + th / 2 + pad_y), radius=14, fill="white", outline="#B8C5D5", width=4)
    draw.text((lx - tw / 2, ly - th / 2 - 4), label, font=LABEL_FONT, fill=INK)


def draw_entity(draw, name, pos):
    x, y = pos
    fill, border = PALETTE[ENTITY_STYLE[name]]
    draw.rounded_rectangle((x, y, x + BOX_W, y + BOX_H), radius=25, fill="white", outline=border, width=8)
    draw.rounded_rectangle((x, y, x + BOX_W, y + 110), radius=25, fill=fill, outline=border, width=8)
    draw.rectangle((x + 4, y + 82, x + BOX_W - 4, y + 112), fill=fill)
    hb = draw.textbbox((0, 0), name, font=HEADER_FONT)
    draw.text((x + (BOX_W - (hb[2] - hb[0])) / 2, y + 24), name, font=HEADER_FONT, fill=INK)
    fields = FIELDS[name]
    gap = min(64, 500 // max(1, len(fields)))
    yy = y + 132
    for field in fields:
        key = field.startswith(("PK", "FK", "UK"))
        draw.text((x + 40, yy), field, font=KEY_FONT if key else BODY_FONT, fill=INK if key else MUTED)
        yy += gap


def make_diagram(filename, positions, edges):
    image = Image.new("RGB", (W, H), BG)
    draw = ImageDraw.Draw(image)
    for parent, child, label, route in edges:
        draw_relationship(draw, positions, parent, child, label, route)
    for name, pos in positions.items():
        draw_entity(draw, name, pos)
    legend = "Relationship label format: parent cardinality   relationship   child cardinality"
    draw.text((90, 2635), legend, font=get_font(34), fill=MUTED)
    path = BUILD / filename
    image.save(path, dpi=(420, 420), quality=95)
    return path


def diagrams():
    results = []
    p = {"ACADEMIC_YEARS": (1860, 80), "SEMESTERS": (540, 930), "CLASSES": (3180, 930), "SUBJECTS": (540, 1840), "STUDENTS": (3180, 1840)}
    e = [
        ("ACADEMIC_YEARS", "SEMESTERS", "1  contains  0..*", None),
        ("ACADEMIC_YEARS", "CLASSES", "1  contains  0..*", None),
        ("SEMESTERS", "SUBJECTS", "0..1  contains  0..*", None),
        ("SEMESTERS", "STUDENTS", "0..1  assigns  0..*", [(1600, 1500), (3000, 1500)]),
        ("CLASSES", "STUDENTS", "0..1  contains  0..*", None),
    ]
    results.append(make_diagram("01_academic.png", p, e))

    p = {"USERS": (1860, 70), "SUBJECTS": (3500, 70), "STUDENTS": (350, 1000), "TEACHERS": (3300, 1000), "API_TOKENS": (350, 1940), "NOTIFICATIONS": (1860, 1940)}
    e = [
        ("USERS", "STUDENTS", "1  has  0..1", None),
        ("USERS", "TEACHERS", "1  has  0..1", None),
        ("USERS", "API_TOKENS", "1  owns  0..*", [(1450, 600), (120, 600), (120, 2250)]),
        ("USERS", "NOTIFICATIONS", "1  receives  0..*", None),
        ("SUBJECTS", "TEACHERS", "0..1  primary for  0..*", None),
    ]
    results.append(make_diagram("02_users.png", p, e))

    p = {"TEACHERS": (180, 80), "SEMESTERS": (1860, 80), "CLASSES": (3540, 80), "TEACHER_TERMS": (1860, 1020), "SUBJECTS": (350, 1940), "TEACHER_SUBJECTS": (2800, 1940)}
    e = [
        ("TEACHERS", "TEACHER_TERMS", "1  assigned  0..*", None),
        ("SEMESTERS", "TEACHER_TERMS", "1  includes  0..*", None),
        ("CLASSES", "TEACHER_TERMS", "1  includes  0..*", None),
        ("TEACHERS", "TEACHER_SUBJECTS", "1  teaches  0..*", [(720, 820), (4550, 820), (4550, 2250)]),
        ("TEACHER_TERMS", "TEACHER_SUBJECTS", "0..1  contains  0..*", None),
        ("SUBJECTS", "TEACHER_SUBJECTS", "1  assigned in  0..*", None),
    ]
    results.append(make_diagram("03_assignments.png", p, e))

    p = {"TEACHERS": (120, 80), "SUBJECTS": (1860, 80), "CLASSROOMS": (3600, 80), "SCHEDULES": (900, 1050), "TEACHER_SUBJECTS": (3000, 1050), "ATTENDANCE_SESSIONS": (1950, 1960)}
    e = [
        ("TEACHERS", "SCHEDULES", "1  follows  0..*", None),
        ("SUBJECTS", "SCHEDULES", "1  scheduled as  0..*", None),
        ("CLASSROOMS", "SCHEDULES", "0..1  hosts  0..*", [(4140, 820), (1450, 820)]),
        ("TEACHERS", "ATTENDANCE_SESSIONS", "1  creates  0..*", [(80, 820), (80, 2450), (1850, 2450)]),
        ("SUBJECTS", "ATTENDANCE_SESSIONS", "1  used by  0..*", None),
        ("TEACHER_SUBJECTS", "ATTENDANCE_SESSIONS", "0..1  authorizes  0..*", None),
        ("SCHEDULES", "ATTENDANCE_SESSIONS", "0..1  starts  0..*", None),
    ]
    results.append(make_diagram("04_scheduling.png", p, e))

    p = {"ATTENDANCE_SESSIONS": (1860, 120), "STUDENTS": (220, 120), "QR_CODES": (900, 1740), "ATTENDANCE": (3000, 1740)}
    e = [
        ("ATTENDANCE_SESSIONS", "QR_CODES", "1  generates  0..*", None),
        ("ATTENDANCE_SESSIONS", "ATTENDANCE", "1  contains  0..*", None),
        ("STUDENTS", "ATTENDANCE", "1  submits  0..*", [(760, 1050), (4300, 1050)]),
    ]
    results.append(make_diagram("05_attendance.png", p, e))
    return results


def remove_borders(paragraph_or_style):
    ppr = paragraph_or_style._element.get_or_add_pPr()
    p_bdr = ppr.find(qn("w:pBdr"))
    if p_bdr is not None:
        ppr.remove(p_bdr)


def shade_cell(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_borders(cell, color="D9D9D9"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        item = borders.find(tag)
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "6")
        item.set(qn("w:color"), color)


def set_cell_margins(cell, value=90):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for edge in ("top", "start", "bottom", "end"):
        node = mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    trpr.append(cant_split)


def add_diagram_page(doc, heading, description, image_path, first=False):
    if not first:
        doc.add_page_break()
    style = "Title" if first else "Heading 1"
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.add_run(heading)
    remove_borders(p)
    intro = doc.add_paragraph(description)
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_after = Pt(7)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_after = Pt(0)
    pic.add_run().add_picture(str(image_path), width=Inches(10.55))


def build_docx(images):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(13)
    section.bottom_margin = Mm(10)
    section.left_margin = Mm(11)
    section.right_margin = Mm(11)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    remove_borders(title)
    heading = doc.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(17)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    remove_borders(heading)

    pages = [
        ("EasyAttend Academic Structure Relationships", "Academic years organize semesters and classes. Semesters organize subjects, while students are assigned to a semester and class.", images[0]),
        ("User Profile and Support Relationships", "Users provide the shared login identity for student and teacher profiles. The same user record owns authentication tokens and receives notifications.", images[1]),
        ("Teaching Assignment Relationships", "Teacher term records connect a teacher to a semester and class. Teacher subject records then connect that term assignment to the subjects being taught.", images[2]),
        ("Scheduling and Attendance Session Relationships", "Schedules connect teachers, subjects, and classrooms. Attendance sessions are created by teachers and may reference a schedule and an authorized teacher subject assignment.", images[3]),
        ("QR GPS and Student Attendance Relationships", "An attendance session generates QR codes and attendance records. Each attendance record belongs to one student and stores GPS coordinates, accuracy, and distance from the configured classroom area.", images[4]),
    ]
    for index, (heading_text, description, image_path) in enumerate(pages):
        add_diagram_page(doc, heading_text, description, image_path, first=index == 0)

    doc.add_page_break()
    p = doc.add_paragraph("Complete Database Relationship Reference", style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_borders(p)
    intro = doc.add_paragraph("This reference lists every foreign key relationship represented across the five diagrams. Cardinality is shown from the parent entity to the child entity.")
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_after = Pt(8)

    table = doc.add_table(rows=1, cols=5)
    table.autofit = False
    widths = [1.55, 0.9, 1.75, 2.15, 4.0]
    headers = ["Parent entity", "Cardinality", "Child entity", "Foreign key", "Relationship meaning"]
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, "1F4E78")
        set_cell_borders(cell)
        set_cell_margins(cell, 100)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
    repeat_table_header(table.rows[0])

    for row_index, (parent, parent_card, child, child_card, foreign_key, meaning) in enumerate(RELATIONSHIPS):
        new_row = table.add_row()
        prevent_row_split(new_row)
        cells = new_row.cells
        values = [parent, f"{parent_card} to {child_card}", child, foreign_key, meaning]
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, "EDF4FA" if row_index % 2 else "FFFFFF")
            set_cell_borders(cell)
            set_cell_margins(cell, 85)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 1 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(8.5)

    props = doc.core_properties
    props.title = "EasyAttend Database Entity Relationship Diagrams"
    props.subject = "Clear A4 database entity relationship diagrams"
    props.author = "EasyAttend Project Team"
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    image_paths = diagrams()
    build_docx(image_paths)
    print(DOCX_PATH)
