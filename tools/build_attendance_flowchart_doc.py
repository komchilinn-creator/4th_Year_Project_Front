from pathlib import Path
from math import atan2, cos, sin, pi

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "tmp" / "attendance_flowchart_document"
BUILD.mkdir(parents=True, exist_ok=True)
DOCX_PATH = BUILD / "EasyAttend_QR_GPS_Attendance_Flowchart_A4.docx"

W, H = 3300, 3600
BG = "#FFFFFF"
INK = "#17315C"
ARROW = "#718096"


def font(size, bold=False):
    names = ["arialbd.ttf", "calibrib.ttf"] if bold else ["arial.ttf", "calibri.ttf"]
    for name in names:
        p = Path("C:/Windows/Fonts") / name
        if p.exists():
            return ImageFont.truetype(str(p), size=size)
    return ImageFont.load_default()


NODE_FONT = font(55, True)
LABEL_FONT = font(43, True)


def multiline_center(draw, box, text, text_font=NODE_FONT, fill=INK, spacing=8):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=text_font, spacing=spacing, align="center")
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2 - 4), text, font=text_font, fill=fill, spacing=spacing, align="center")


def process(draw, center, text, kind="process", width=1500, height=220):
    cx, cy = center
    box = (cx - width / 2, cy - height / 2, cx + width / 2, cy + height / 2)
    styles = {
        "process": ("#D9EAF7", "#3D85C6"),
        "error": ("#F4CCCC", "#CC0000"),
        "success": ("#D9D2E9", "#674EA7"),
        "connector": ("#EDF2F7", "#64748B"),
    }
    fill, stroke = styles[kind]
    draw.rounded_rectangle(box, radius=40, fill=fill, outline=stroke, width=10)
    multiline_center(draw, box, text)
    return box


def terminator(draw, center, text, width=1000, height=210):
    cx, cy = center
    box = (cx - width / 2, cy - height / 2, cx + width / 2, cy + height / 2)
    draw.rounded_rectangle(box, radius=height / 2, fill="#D9EAD3", outline="#38761D", width=11)
    multiline_center(draw, box, text)
    return box


def decision(draw, center, text, width=1700, height=420):
    cx, cy = center
    pts = [(cx, cy - height / 2), (cx + width / 2, cy), (cx, cy + height / 2), (cx - width / 2, cy)]
    draw.polygon(pts, fill="#FFF2CC", outline="#BF9000")
    draw.line(pts + [pts[0]], fill="#BF9000", width=11, joint="curve")
    multiline_center(draw, (cx - width * .32, cy - height * .30, cx + width * .32, cy + height * .30), text)
    return (cx - width / 2, cy - height / 2, cx + width / 2, cy + height / 2)


def arrow_head(draw, previous, end):
    angle = atan2(end[1] - previous[1], end[0] - previous[0])
    length = 35
    spread = pi / 7
    a = (end[0] - length * cos(angle - spread), end[1] - length * sin(angle - spread))
    b = (end[0] - length * cos(angle + spread), end[1] - length * sin(angle + spread))
    draw.polygon([end, a, b], fill=ARROW)


def connect(draw, points, label=None, label_at=None):
    draw.line(points, fill=ARROW, width=10, joint="curve")
    arrow_head(draw, points[-2], points[-1])
    if label:
        lx, ly = label_at or points[len(points) // 2]
        bbox = draw.textbbox((0, 0), label, font=LABEL_FONT)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.rounded_rectangle((lx - tw / 2 - 18, ly - th / 2 - 10, lx + tw / 2 + 18, ly + th / 2 + 10), radius=18, fill="white", outline="#CBD5E1", width=4)
        draw.text((lx - tw / 2, ly - th / 2 - 4), label, font=LABEL_FONT, fill=INK)


def page_one():
    image = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(image)
    boxes = {}
    boxes["A"] = terminator(d, (1650, 150), "Start")
    boxes["B"] = process(d, (1650, 500), "Teacher Creates\nAttendance Session")
    boxes["C"] = process(d, (1650, 850), "System Generates QR Code")
    boxes["D"] = process(d, (1650, 1200), "Student Opens QR Scanner")
    boxes["E"] = process(d, (1650, 1550), "Student Scans QR Code")
    boxes["F"] = process(d, (1650, 1900), "Request GPS Location")
    boxes["G"] = decision(d, (1650, 2400), "GPS Available and\nPermission Granted?")
    boxes["I"] = process(d, (1450, 2900), "Get Current Location", width=1200)
    boxes["Q"] = process(d, (1450, 3330), "Continue to QR Validation", "connector", width=1200)
    boxes["H"] = process(d, (2700, 2900), "Show GPS Error", "error", width=900)
    boxes["P"] = terminator(d, (2700, 3330), "End", width=700)

    sequence = ["A", "B", "C", "D", "E", "F"]
    for first, second in zip(sequence, sequence[1:]):
        connect(d, [(1650, boxes[first][3]), (1650, boxes[second][1])])
    connect(d, [(1650, boxes["F"][3]), (1650, boxes["G"][1])])
    connect(d, [(1650, boxes["G"][3]), (1650, 2680), (1450, 2680), (1450, boxes["I"][1])], "Yes", (1510, 2640))
    connect(d, [(1450, boxes["I"][3]), (1450, boxes["Q"][1])])
    connect(d, [(boxes["G"][2], 2400), (2700, 2400), (2700, boxes["H"][1])], "No", (2600, 2530))
    connect(d, [(2700, boxes["H"][3]), (2700, boxes["P"][1])])
    path = BUILD / "attendance_flow_part1.png"
    image.save(path, dpi=(420, 420), quality=95)
    return path


def page_two():
    image = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(image)
    boxes = {}
    boxes["Q"] = process(d, (1650, 180), "Location Received", "connector", width=1200)
    boxes["J"] = decision(d, (1650, 650), "QR Code Valid?", width=1450)
    boxes["L"] = decision(d, (1650, 1350), "Location Accurate and\nInside Allowed Area?")
    boxes["N"] = process(d, (1650, 1950), "Record Student Attendance")
    boxes["O"] = process(d, (1650, 2400), "Show Attendance Success", "success")
    boxes["P"] = terminator(d, (1650, 3100), "End")
    boxes["K"] = process(d, (2700, 950), "Show QR Error", "error", width=1050)
    boxes["M"] = process(d, (2700, 1700), "Show Location Error", "error", width=1200)

    connect(d, [(1650, boxes["Q"][3]), (1650, boxes["J"][1])])
    connect(d, [(1650, boxes["J"][3]), (1650, boxes["L"][1])], "Yes", (1530, 1000))
    connect(d, [(boxes["J"][2], 650), (2700, 650), (2700, boxes["K"][1])], "No", (2600, 760))
    connect(d, [(1650, boxes["L"][3]), (1650, boxes["N"][1])], "Yes", (1530, 1650))
    connect(d, [(boxes["L"][2], 1350), (2700, 1350), (2700, boxes["M"][1])], "No", (2600, 1470))
    connect(d, [(1650, boxes["N"][3]), (1650, boxes["O"][1])])
    connect(d, [(1650, boxes["O"][3]), (1650, boxes["P"][1])])
    connect(d, [(2700, boxes["K"][3]), (3100, boxes["K"][3]), (3100, 2850), (1900, 2850), (1900, boxes["P"][1])])
    connect(d, [(2700, boxes["M"][3]), (3000, boxes["M"][3]), (3000, 2750), (1800, 2750), (1800, boxes["P"][1])])
    path = BUILD / "attendance_flow_part2.png"
    image.save(path, dpi=(420, 420), quality=95)
    return path


def remove_borders(paragraph_or_style):
    ppr = paragraph_or_style._element.get_or_add_pPr()
    p_bdr = ppr.find(qn("w:pBdr"))
    if p_bdr is not None:
        ppr.remove(p_bdr)


def add_page(doc, heading, description, image_path, first=False):
    if not first:
        doc.add_page_break()
    p = doc.add_paragraph(heading, style="Title" if first else "Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    remove_borders(p)
    intro = doc.add_paragraph(description)
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.paragraph_format.space_after = Pt(5)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_after = Pt(0)
    pic.add_run().add_picture(str(image_path), width=Inches(6.95))


def build_docx(first_image, second_image):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(12)
    section.bottom_margin = Mm(10)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(19)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    remove_borders(title)
    heading = doc.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(17)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    remove_borders(heading)

    add_page(
        doc,
        "EasyAttend QR and GPS Attendance Flowchart",
        "Part 1 shows attendance-session creation, QR scanning, and the browser GPS permission check.",
        first_image,
        first=True,
    )
    add_page(
        doc,
        "QR Location Validation and Attendance Recording",
        "Part 2 shows QR validation, permitted-area validation, attendance recording, and each possible error outcome.",
        second_image,
    )
    props = doc.core_properties
    props.title = "EasyAttend QR and GPS Attendance Flowchart"
    props.subject = "A4 top down attendance flowchart"
    props.author = "EasyAttend Project Team"
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    p1 = page_one()
    p2 = page_two()
    build_docx(p1, p2)
    print(DOCX_PATH)
