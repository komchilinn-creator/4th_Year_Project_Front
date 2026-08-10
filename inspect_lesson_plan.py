from docx import Document

path = r"D:\WYTU\4th_Year_Project\ICT_Year5_Unit3_2nd_Term_lesson_plan_expanded.docx"
doc = Document(path)

print(f"PARAGRAPHS: {len(doc.paragraphs)}")
for i, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if text:
        print(f"P {i}: {text}")

print(f"TABLES: {len(doc.tables)}")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti}: {len(table.rows)} rows x {len(table.columns)} cols")
    for ri, row in enumerate(table.rows):
        cells = [cell.text.replace("\n", " | ").strip() for cell in row.cells]
        print(f"R {ri}: {cells}")
