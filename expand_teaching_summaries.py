from docx import Document

path = r"D:\WYTU\4th_Year_Project\ICT_Year5_Unit3_2nd_Term_lesson_plan_expanded.docx"

summaries = [
    "This class will check what students already know about computers, networks, the internet, and online safety. Students will answer simple questions and share their ideas with the class. This will help the teacher see what students understand already and what they need to learn more about before we start the unit.",
    "In this class, students will learn what a network is and why we use it. They will see how different devices, such as computers, phones, printers, and tablets, can connect and share information. Then, they will learn about wired and wireless networks and look at simple examples from their school, home, and daily life.",
    "This class will show students how the internet really works. We will start from small networks to big networks, including wired and wireless. We will look at the infrastructure, the physical parts like the huge cables under the ocean that connect the whole world, and see how it all links to their own phone by wireless. After the class, students will make a concept map to show their understanding.",
    "Students will learn how to draw a simple networking diagram. They will start with the main devices, such as a computer, server, switch, router, and hub, and learn the symbols used for each one. Then, they will see how everything connects in a simple network diagram and make their own diagram for a classroom, home, or school network.",
    "In this class, students will learn about network protocols, especially Transmission Control Protocol (TCP) and Internet Protocol (IP). They will learn how TCP works to ensure data arrives correctly and how IP routes the data to the right address. Finally, they will build a complete model network using the TCP/IP protocol to see how it all works together in practice.",
    "This class will help students understand the difference between the Internet and the World Wide Web. They will learn how a web browser, web server, URL, and HTTP work together when they open a website. Then, they will follow the journey of a webpage from a browser to a server and back to their computer.",
    "This class expands on websites by showing how a website is organized and how people move from one page to another. Students will learn about home pages, menus, links, images, headings, and contact pages. They will also look at a website and see how the front-end that users see connects to the back-end that helps the website work.",
    "In this class, students will learn about social engineering and how hackers can use tricks to get information from people. They will look at examples of urgent messages, fake prizes, and requests for personal details. Finally, they will learn that they should stop, check with a trusted adult, and never share private information with strangers online.",
    "This class will help students learn about phishing and pharming. They will see how fake emails, messages, and websites try to trick people into giving passwords or personal information. Students will practice checking the sender, spelling, links, and website address, and learn how to stay safe by not clicking unknown links and asking a trusted adult for help.",
    "In this class, students will learn how to stay secure when they use the internet. They will learn about strong passwords, software updates, privacy settings, and logging out from shared computers. Finally, they will make a simple security checklist to show the safe things they can do at school and at home.",
    "This class will introduce biometric security, such as fingerprint, face, and voice recognition. Students will learn how these methods can help protect a device and why biometric information is important. Finally, they will review the whole unit, including networks, the internet, online threats, and the safe habits they should use every day.",
]

doc = Document(path)
summary_index = 0
for table in doc.tables:
    for row in table.rows:
        if row.cells[0].text.strip().lower() == "teaching summary":
            cell = row.cells[1]
            paragraph = cell.paragraphs[0]
            if paragraph.runs:
                paragraph.runs[0].text = summaries[summary_index]
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(summaries[summary_index])
            summary_index += 1

if summary_index != len(summaries):
    raise RuntimeError(f"Updated {summary_index} summaries; expected {len(summaries)}")

doc.save(path)
print(f"Updated {summary_index} teaching summaries: {path}")
